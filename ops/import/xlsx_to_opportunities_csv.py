#!/usr/bin/env python3
"""Turn the band's booking sheet into one CSV the opportunity importer can read.

`viryaos_team_opportunities` is what the `booking_opportunity` and
`live_opportunity` autopilot contexts rank and act on, and in production it held
zero rows — the whole live arm of the growth loop had never had a unit to work
with. Meanwhile the CRM's `Zgloszenia` sheet holds 797 festivals, competitions
and showcases gathered by hand, with fit scores, contact addresses and deadlines.

Shows are the top of the fan funnel. This is the widest gap between what the band
knows and what the system knows.

Reading and mapping happen here; `crowdrelay-worker import-opportunities` owns the
writes, so imported rows land through the same table and the same status rules as
anything the brain proposes.

Usage:
    python3 ops/import/xlsx_to_opportunities_csv.py --out opportunities.csv
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

WORKBOOK = "VIRYA — CRM _ KONTAKTY — EDYTUJ TU.xlsx"
SHEET = "Zgloszenia"
SOURCE = "crm:zgloszenia"

# `viryaos_team_opportunities_opportunity_kind_check` accepts exactly these five.
# `support_slot` and `funding` are not produced here: nothing in this sheet is a
# support slot or a grant, and guessing one would put a row in front of the
# funding autopilot that has no money in it.
#
# Matched on the first word of the sheet's `Typ`, because the real values are
# compounds — "Festiwal / konkurs", "Konkurs / przegląd", "Festiwal / konkurs;
# Festiwal" — and the first token is the outcome. "Festiwal / konkurs" is a
# festival with a competition entry: playing it is the point, so `festival`.
KIND_BY_FIRST_WORD = {
    "festiwal": "festival",
    "juwenalia": "festival",  # Polish student festivals, a real live circuit
    "konkurs": "review_contest",
    "przeglad": "review_contest",
    "targi": "showcase",
    "showcase": "showcase",
}

COUNTRY = {
    "polska": "PL",
    "niemcy": "DE",
    "czechy": "CZ",
    "uk": "GB",
    "wielka brytania": "GB",
    "anglia": "GB",
    "slowacja": "SK",
    "litwa": "LT",
    "ukraina": "UA",
    "austria": "AT",
    "holandia": "NL",
    "belgia": "BE",
    "francja": "FR",
    "wegry": "HU",
    "szwecja": "SE",
    "dania": "DK",
    "norwegia": "NO",
    "finlandia": "FI",
    "hiszpania": "ES",
    "wlochy": "IT",
    "szwajcaria": "CH",
}

# How much to trust the row itself, as basis points. This is confidence in the
# *record*, not fit: an officially verified 2026 entry and a historical
# ROCK METAL scrape are not the same claim, and the brain should not weigh them
# the same. The sheet says which is which, so this reads it rather than guessing.
CONFIDENCE = {
    "zweryfikowane 2026 - zrodlo oficjalne": 7000,
    "kontakt biezacy - do weryfikacji": 4000,
    "nowszy wpis - do weryfikacji": 4000,
    "historyczny - do weryfikacji": 2500,
    "zrodlo historyczne rock metal - do podwojnej weryfikacji": 2000,
}
DEFAULT_CONFIDENCE = 2500


def fold(value: object) -> str:
    """Lowercase, strip Polish diacritics and collapse dashes for lookups."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    # Every Polish diacritic decomposes under NFKD except ł, which is a distinct
    # letter with no canonical decomposition — so the loop above leaves it and a
    # lookup key written "spolecznosc" never matches "społeczność". That silently
    # dropped whole categories: the booking sheet's "Dokładny — edycja 2026"
    # deadline type matched nothing, so 23 exact festival deadlines were read as
    # estimates. Mapped by hand, both cases, before any lookup.
    text = text.replace("ł", "l").replace("Ł", "L")
    text = text.replace("—", "-").replace("–", "-")
    return re.sub(r"\s+", " ", text).strip().lower()


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def first_email(value: object) -> str:
    """The first address in a cell. Several rows carry two, comma separated."""
    match = re.search(r"[^\s,;<>()]+@[^\s,;<>()]+\.[^\s,;<>()]+", str(value or ""))
    return match.group(0).strip(".,;") if match else ""


def first_url(value: object) -> str:
    text = clean(value)
    return text if text.startswith(("http://", "https://")) else ""


def basis_points(value: object) -> str:
    """`Dopasowanie_Virya` is 0-100; the column is 0-10000."""
    try:
        score = float(value)
    except (TypeError, ValueError):
        return ""
    return str(max(0, min(10000, round(score * 100))))


def exact_deadline(raw: object, deadline_type: object) -> str:
    """An ISO date, and only when the sheet calls the deadline exact.

    429 rows carry `Szacunek z miesiąca wydarzenia` — an estimate derived from
    the month the event happens in — and 278 carry no deadline at all. Writing
    either into `deadline` would put an invented date in the column the brain
    sorts its live calendar by, and the brain cannot tell an estimate from a
    fact once it is a timestamp. The month estimate travels in `metadata`
    instead, where it reads as what it is.
    """
    if not fold(deadline_type).startswith("dokladny"):
        return ""
    text = clean(raw)
    for parse in (
        lambda t: datetime.fromisoformat(t).date(),
        lambda t: datetime.strptime(t, "%d.%m.%Y").date(),
        lambda t: datetime.strptime(t, "%Y-%m-%d").date(),
    ):
        try:
            parsed = parse(text)
        except (ValueError, TypeError):
            continue
        return parsed.isoformat() if isinstance(parsed, date) else ""
    return ""


def load_rows(path: Path, sheet: str) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        print(
            "openpyxl is required: pip install --user openpyxl",
            file=sys.stderr,
        )
        raise SystemExit(2) from None
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in book.sheetnames:
        raise SystemExit(f"{path.name} has no sheet {sheet!r}")
    worksheet = book[sheet]
    rows = worksheet.iter_rows(values_only=True)
    header = [clean(cell) for cell in next(rows, ())]
    records = []
    for row in rows:
        record = dict(zip(header, row, strict=False))
        if clean(record.get("Wydarzenie")):
            records.append(record)
    book.close()
    return records


FIELDS = [
    "external_key",
    "opportunity_kind",
    "title",
    "organization",
    "contact_email",
    "destination_url",
    "country_code",
    "city",
    "fit_basis_points",
    "confidence_basis_points",
    "deadline",
    "eligible",
    "typical_month",
    "deadline_certainty",
    "priority",
    "genres",
    "next_step",
    "why_fit",
    "verification",
]


def convert(records: list[dict]) -> tuple[list[dict], dict[str, int]]:
    out: list[dict] = []
    counts = {
        "read": len(records),
        "no_route": 0,
        "unmapped_kind": 0,
        "confirmed_inactive": 0,
        "not_eligible": 0,
        "with_deadline": 0,
    }
    for record in records:
        title = clean(record.get("Wydarzenie"))
        activity = fold(record.get("Status_aktywności"))
        # A festival somebody already confirmed is gone is not an opportunity.
        # Importing it would spend a dispatch to reach a dead inbox and teach the
        # brain that festival outreach does not work.
        if activity.startswith("nieaktywne - potwierdzone"):
            counts["confirmed_inactive"] += 1
            continue

        kind = KIND_BY_FIRST_WORD.get(fold(record.get("Typ")).split(" ")[0].split("/")[0])
        if not kind:
            counts["unmapped_kind"] += 1
            continue

        email = first_email(record.get("Email_kontaktowy"))
        form = first_url(record.get("Formularz_URL"))
        site = first_url(record.get("Strona_WWW")) or first_url(record.get("WWW"))
        # A row with neither an address nor a submission form has no route. The
        # press-pitch lesson: an action with no destination is a succeeded
        # dispatch that reached nobody, and it is worse than no row at all.
        if not email and not form:
            counts["no_route"] += 1
            continue

        openness = fold(record.get("Zgłoszenia_otwarte"))
        eligible = not (
            openness.startswith("zamkniete") or openness == "nie"
        )
        if not eligible:
            counts["not_eligible"] += 1

        deadline = exact_deadline(
            record.get("Deadline_aktualny"), record.get("Deadline_typ")
        )
        if deadline:
            counts["with_deadline"] += 1

        month = clean(record.get("Deadline_miesiac_typowy"))
        if month.endswith(".0"):
            month = month[:-2]

        out.append(
            {
                "external_key": clean(record.get("Zgłoszenie_ID")) or fold(title)[:200],
                "opportunity_kind": kind,
                "title": title[:240],
                # The sheet has no separate organiser column: for a festival the
                # event and the organisation are the same party.
                "organization": title[:240],
                "contact_email": email[:320],
                "destination_url": (form or site)[:2048],
                "country_code": COUNTRY.get(fold(record.get("Kraj")), ""),
                "city": clean(record.get("Miasto"))[:120],
                "fit_basis_points": basis_points(record.get("Dopasowanie_Virya")),
                "confidence_basis_points": str(
                    CONFIDENCE.get(fold(record.get("Status_weryfikacji")), DEFAULT_CONFIDENCE)
                ),
                "deadline": deadline,
                "eligible": "true" if eligible else "false",
                "typical_month": month,
                "deadline_certainty": clean(record.get("Deadline_pewnosc")),
                "priority": clean(record.get("Priorytet")),
                "genres": clean(record.get("Gatunki"))[:300],
                "next_step": clean(record.get("Następny_krok"))[:500],
                "why_fit": clean(record.get("Uzasadnienie_dopasowania"))[:1000],
                "verification": clean(record.get("Status_weryfikacji"))[:200],
            }
        )
    return out, counts


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", default=str(Path.home() / "dev" / "contacts"))
    parser.add_argument("--out", default="opportunities.csv")
    args = parser.parse_args()

    path = Path(args.source_dir) / WORKBOOK
    if not path.is_file():
        print(f"missing workbook: {path}", file=sys.stderr)
        return 2

    rows, counts = convert(load_rows(path, SHEET))
    # Dedupe on the key the table is unique on, keeping the first occurrence.
    seen: set[str] = set()
    unique = []
    for row in rows:
        if row["external_key"] in seen:
            continue
        seen.add(row["external_key"])
        unique.append(row)
    counts["duplicate_key"] = len(rows) - len(unique)

    out = Path(args.out)
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(unique)

    print(f"wrote {len(unique)} opportunities to {out}")
    for key in sorted(counts):
        print(f"  {key}: {counts[key]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
